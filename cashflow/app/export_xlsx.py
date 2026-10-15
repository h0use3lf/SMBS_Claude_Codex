"""부속명세서 Excel 생성 (수식 포함)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import core

FONT = "맑은 고딕"
THIN = Side(style="thin", color="000000")
HAIR = Side(style="thin", color="A6A6A6")
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
NUM = '#,##0;[Red]-#,##0;"-"'
DIFF = '+#,##0;[Red]-#,##0;"0"'


def _f(bold=False, size=9):
    return Font(name=FONT, size=size, bold=bold)


def _setup(ws, widths, title, sub):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["A1"] = title
    ws["A1"].font = _f(True, 13)
    ws["A2"] = sub
    ws["A2"].font = _f(False, 9)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.font = FONT
    ws.sheet_view.showGridLines = False


def _header(ws, row, heads):
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = _f(True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=THIN, bottom=THIN, left=HAIR, right=HAIR)
    ws.row_dimensions[row].height = 20
    ws.print_title_rows = f"{row}:{row}"


def _cell(ws, r, c, v, fmt=None, bold=False, align=None, fill=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = _f(bold)
    x.border = Border(top=HAIR, bottom=HAIR, left=HAIR, right=HAIR)
    if fmt:
        x.number_format = fmt
        x.alignment = Alignment(horizontal="right", vertical="center")
    else:
        x.alignment = Alignment(horizontal=align or "left", vertical="center", wrap_text=True)
    if fill:
        x.fill = fill
    return x


def _close(ws, r, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = Border(top=cell.border.top, bottom=THIN, left=HAIR, right=HAIR)


def sheet_budget(wb, data):
    ws = wb.create_sheet("1.예산표")
    _setup(ws, [11, 20, 13, 13, 13, 34, 40], "부속명세서 1. 월 예산 및 2주 기간 예산",
           f"예산안: {core.TIERS.get(data['budget'].get('choice') or 'B')['name']} │ 단위: 원")
    heads = ["구분", "항목", "한 달 예산", "상반기", "후반기", "설정 근거", "2주 설정 근거"]
    _header(ws, 4, heads)
    r = 5
    rows = core.budget_rows(data)
    first = r
    for g in ["소득", "저축", "고정지출", "변동지출", "예비비"]:
        gr = [x for x in rows if x[0] == g]
        if not gr:
            continue
        s = r
        for x in gr:
            _cell(ws, r, 1, g if r == s else None, align="center")
            _cell(ws, r, 2, x[1])
            for i, v in enumerate(x[2:5], 3):
                _cell(ws, r, i, v, NUM)
            _cell(ws, r, 6, x[5])
            _cell(ws, r, 7, x[6])
            r += 1
        _cell(ws, r, 1, None)
        _cell(ws, r, 2, f"{g} 소계", bold=True, fill=SUB_FILL)
        for i in (3, 4, 5):
            L = get_column_letter(i)
            _cell(ws, r, i, f"=SUM({L}{s}:{L}{r - 1})", NUM, bold=True, fill=SUB_FILL)
        _cell(ws, r, 6, None, fill=SUB_FILL)
        _cell(ws, r, 7, None, fill=SUB_FILL)
        if r - s > 1:
            ws.merge_cells(start_row=s, start_column=1, end_row=r - 1, end_column=1)
        _close(ws, r, 7)
        r += 1
    r += 1
    _cell(ws, r, 2, "예산 검증(소득−저축−고정−변동−예비비)", bold=True)
    for i in (3, 4, 5):
        L = get_column_letter(i)
        f = (f'=SUMIF($B${first}:$B${r-2},"소득 소계",{L}{first}:{L}{r-2})'
             f'-SUMIF($B${first}:$B${r-2},"저축 소계",{L}{first}:{L}{r-2})'
             f'-SUMIF($B${first}:$B${r-2},"고정지출 소계",{L}{first}:{L}{r-2})'
             f'-SUMIF($B${first}:$B${r-2},"변동지출 소계",{L}{first}:{L}{r-2})'
             f'-SUMIF($B${first}:$B${r-2},"예비비 소계",{L}{first}:{L}{r-2})')
        _cell(ws, r, i, f, NUM, bold=True)
    _cell(ws, r + 1, 2, "※ 양(+)이면 해당 기간 말 이월 예정액, 음(−)이면 이월 잔액 사용 예정액")
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=7)


def sheet_settle(wb, data):
    ws = wb.create_sheet("2.예산대비실적")
    _setup(ws, [11, 20, 13, 13, 13, 48], "부속명세서 2. 기간 예산 대비 실적",
           "차이 = 실적 − 예산 │ 소득은 (+)가 유리, 지출·저축은 (+)가 초과 │ 단위: 원")
    r = 4
    for key in ("p1", "p2", "month"):
        an = core.analyze(data, key)
        ws.cell(row=r, column=1, value=an["title"]).font = _f(True, 10)
        r += 1
        _header(ws, r, ["구분", "항목", "예산", "실적", "차이", "비고"])
        r += 1
        s = r
        last = None
        for g, n, b, a in an["rows"]:
            _cell(ws, r, 1, g if g != last else None, align="center")
            last = g
            _cell(ws, r, 2, n)
            _cell(ws, r, 3, b, NUM)
            _cell(ws, r, 4, a, NUM)
            _cell(ws, r, 5, f"=D{r}-C{r}", DIFF)
            _cell(ws, r, 6, core.remark(data, an, g, n))
            r += 1
        e = r - 1
        rin, rout, rnet = r, r + 1, r + 2
        inc_end = s + sum(1 for x in an["rows"] if x[0] == "소득") - 1
        for rr, label, col_f in [
            (rin, "유입 합계", lambda L: f"=SUM({L}{s}:{L}{inc_end})"),
            (rout, "유출 합계", lambda L: f"=SUM({L}{inc_end+1}:{L}{e})"),
            (rnet, "순현금흐름(유입−유출)", lambda L: f"={L}{rin}-{L}{rout}"),
        ]:
            _cell(ws, rr, 1, None, fill=SUB_FILL)
            _cell(ws, rr, 2, label, bold=True, fill=SUB_FILL)
            _cell(ws, rr, 3, col_f("C"), NUM, True, fill=SUB_FILL)
            _cell(ws, rr, 4, col_f("D"), NUM, True, fill=SUB_FILL)
            _cell(ws, rr, 5, f"=D{rr}-C{rr}", DIFF, True, fill=SUB_FILL)
            _cell(ws, rr, 6, None, fill=SUB_FILL)
        r = rnet + 1
        _cell(ws, r, 1, None)
        _cell(ws, r, 2, "예비비", bold=True)
        _cell(ws, r, 3, an["reserve_b"], NUM)
        _cell(ws, r, 4, an["reserve_used"], NUM)
        _cell(ws, r, 5, f"=C{r}-D{r}", DIFF)
        _cell(ws, r, 6, "실적 = 소비지출 예산 초과분(예비비 사용액), 차이 = 잔여 예비비")
        _close(ws, r, 6)
        r += 3


def sheet_daily(wb, data):
    ws = wb.create_sheet("3.일별잔액")
    _setup(ws, [12, 8, 8, 13, 13, 13, 14, 10], "부속명세서 3. 일별 현금흐름 및 잔액 추적",
           "잔액 = 전일 잔액 + 유입 − 소비 − 저축 │ 단위: 원")
    an = core.analyze(data, "month")
    _header(ws, 4, ["일자", "요일", "구분", "유입", "소비지출", "저축·투자", "일말 잔액(장부)", "기록 건수"])
    _cell(ws, 5, 1, "기초")
    for c in range(2, 7):
        _cell(ws, 5, c, None)
    _cell(ws, 5, 7, an["start"], NUM, True)
    _cell(ws, 5, 8, None)
    yo = "월화수목금토일"
    r = 6
    for x in an["daily"]:
        dt = x["date"]
        _cell(ws, r, 1, dt.strftime("%Y-%m-%d"), align="center")
        _cell(ws, r, 2, yo[dt.weekday()], align="center")
        _cell(ws, r, 3, "휴일" if x["holiday"] else "평일", align="center")
        _cell(ws, r, 4, x["in"], NUM)
        _cell(ws, r, 5, x["exp"], NUM)
        _cell(ws, r, 6, x["sav"], NUM)
        _cell(ws, r, 7, f"=G{r-1}+D{r}-E{r}-F{r}", NUM)
        _cell(ws, r, 8, x["count"] if x["count"] else "지출 없음" if not x["exp"] else 0, align="center")
        if x["holiday"]:
            for c in range(1, 9):
                ws.cell(row=r, column=c).fill = SUB_FILL
        r += 1
    _cell(ws, r, 1, "합계", bold=True)
    for c in (2, 3):
        _cell(ws, r, c, None)
    for c in (4, 5, 6):
        L = get_column_letter(c)
        _cell(ws, r, c, f"=SUM({L}6:{L}{r-1})", NUM, True)
    _cell(ws, r, 7, f"=G{r-1}", NUM, True)
    _cell(ws, r, 8, None)
    _close(ws, r, 8)
    _cell(ws, r + 1, 1, "최저 잔액")
    _cell(ws, r + 1, 7, f"=MIN(G6:G{r-1})", NUM, True)
    _cell(ws, r + 1, 2, an["min_day"].strftime("%m/%d"), align="center")


def sheet_ledger(wb, data):
    ws = wb.create_sheet("4.거래원장")
    _setup(ws, [11, 6, 18, 11, 11, 13, 10, 7, 22, 22, 6, 7, 8], "부속명세서 4. 거래 원장(전체)",
           "만족도: 1(매우 낮음)~5(매우 높음) │ 판단: 유지·조정·중단 │ 증빙: 영수증 첨부 여부")
    _header(ws, 4, ["일자", "종류", "지출 항목", "금액", "카테고리", "세부", "결제수단", "계획",
                    "이유·상황", "감정", "만족", "판단", "증빙"])
    r = 5
    for t in sorted(data["transactions"], key=lambda t: t["date"]):
        vals = [core.d(t["date"]), t["kind"], t["item"], t["amount"], t["category"], t["sub"], t["method"],
                "계획" if t["planned"] else "비계획", t["reason"], t["emotion"],
                t.get("satisfaction", ""), core.judge_of(t), "첨부" if t.get("receipt") else ""]
        for i, v in enumerate(vals, 1):
            x = _cell(ws, r, i, v, NUM if i == 4 else None,
                      align="center" if i in (1, 2, 7, 8, 11, 12, 13) else None)
            if i == 1:
                x.number_format = "yyyy-mm-dd"
        r += 1
    _close(ws, r - 1, 13)
    ws.auto_filter.ref = f"A4:M{max(r-1, 5)}"
    ws.freeze_panes = "A5"
    return r - 1


def sheet_category(wb, data, last):
    ws = wb.create_sheet("5.카테고리분석")
    _setup(ws, [16, 13, 13, 13, 9, 13, 13, 13, 9], "부속명세서 5. 카테고리별 지출 분석",
           "거래원장 수식 집계 │ 비중 = 해당 기간 소비지출 합계 대비")
    ps = core.periods(data)
    heads = ["카테고리"]
    for k in ("p1", "p2"):
        heads += [f"{core.PERIOD_LABEL[k]} 금액", "건수"]
    heads += ["월 합계", "월 비중", "평균 단가", "월 예산", "예산 대비", "계획 비율"]
    _header(ws, 4, heads)
    L = "'4.거래원장'!"
    rng = lambda col: f"{L}${col}$5:${col}${max(last, 5)}"
    cats = ["고정지출"] + core.VAR_CATS
    bud = {x[1]: x[2] for x in core.budget_rows(data) if x[0] == "변동지출"}
    bud["고정지출"] = sum(x[2] for x in core.budget_rows(data) if x[0] == "고정지출")
    r = 5
    for c in cats:
        _cell(ws, r, 1, c)
        col = 2
        for k in ("p1", "p2"):
            a, b = ps[k]
            cond = (f'{rng("B")},"지출",{rng("E")},$A{r},'
                    f'{rng("A")},">="&DATE({a.year},{a.month},{a.day}),'
                    f'{rng("A")},"<="&DATE({b.year},{b.month},{b.day})')
            _cell(ws, r, col, f"=SUMIFS({rng('D')},{cond})", NUM)
            _cell(ws, r, col + 1, f"=COUNTIFS({cond})", "#,##0")
            col += 2
        _cell(ws, r, 6, f"=B{r}+D{r}", NUM)
        _cell(ws, r, 7, f"=IFERROR(F{r}/F${5+len(cats)},0)", "0.0%")
        _cell(ws, r, 8, f"=IFERROR(F{r}/(C{r}+E{r}),0)", NUM)
        _cell(ws, r, 9, bud.get(c, 0), NUM)
        _cell(ws, r, 10, f"=IFERROR(F{r}/I{r},0)", "0.0%")
        _cell(ws, r, 11, f'=IFERROR(SUMIFS({rng("D")},{rng("B")},"지출",{rng("E")},$A{r},{rng("H")},"계획")/F{r},0)',
              "0.0%")
        r += 1
    _cell(ws, r, 1, "합계", bold=True, fill=SUB_FILL)
    for c in range(2, 12):
        Lc = get_column_letter(c)
        if c in (7,):
            _cell(ws, r, c, f"=SUM({Lc}5:{Lc}{r-1})", "0.0%", True, fill=SUB_FILL)
        elif c in (8, 10, 11):
            _cell(ws, r, c, None, fill=SUB_FILL)
        else:
            _cell(ws, r, c, f"=SUM({Lc}5:{Lc}{r-1})", NUM, True, fill=SUB_FILL)
    _close(ws, r, 11)
    ws.column_dimensions["J"].width = 10
    ws.column_dimensions["K"].width = 10


def build(data, path):
    wb = Workbook()
    wb.remove(wb.active)
    sheet_budget(wb, data)
    sheet_settle(wb, data)
    sheet_daily(wb, data)
    last = sheet_ledger(wb, data)
    sheet_category(wb, data, last)
    wb.save(path)
    return path
