"""제출용 Word 보고서 생성."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from . import charts, core

FONT = "맑은 고딕"
W = 170  # 본문 폭(mm)


# ───────────────────────── 서식 도우미 ─────────────────────────
def _font(run, size=9.5, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def _borders(table, outer=12, inner=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge, sz in (("top", outer), ("bottom", outer), ("left", 0), ("right", 0),
                     ("insideH", inner), ("insideV", inner)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if sz else "nil")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), "000000" if edge in ("top", "bottom") else "8C8C8C")
        b.append(e)
    tblPr.append(b)


def _cell_margins(table, mm=1.2):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(mm * 56.7)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def _widths(t, widths):
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths) * 56.7)))
    tw.set(qn("w:type"), "dxa")
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(tw)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 56.7)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Mm(w)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    trPr.append(c)


def _put(cell, text, size=8.5, bold=False, align="left", color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.1
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    _font(p.add_run("" if text is None else str(text)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, heads, rows, widths, num=(), center=(), bold_rows=(), shade_rows=(), merge_first=False,
          size=8.5):
    t = doc.add_table(rows=1, cols=len(heads))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    _cell_margins(t)
    for i, h in enumerate(heads):
        c = t.rows[0].cells[i]
        _put(c, h, size, True, "center")
        _shade(c, "D9D9D9")
    _repeat_header(t.rows[0])
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        _no_split(t.rows[-1])
        for i, v in enumerate(r):
            if i in num and isinstance(v, (int, float)):
                txt, al = core.won(v), "right"
            else:
                txt, al = v, "right" if i in num else "center" if i in center else "left"
            _put(cells[i], txt, size, ri in bold_rows, al)
            if ri in shade_rows:
                _shade(cells[i], "F2F2F2")
    _widths(t, widths)
    if merge_first:  # 연속된 같은 구분값 세로 병합
        ri = 0
        while ri < len(rows):
            v = rows[ri][0]
            rj = ri
            while v not in (None, "") and rj + 1 < len(rows) and rows[rj + 1][0] == v:
                rj += 1
            if rj > ri:
                a = t.cell(ri + 1, 0)
                a.merge(t.cell(rj + 1, 0))
                _put(a, v, size, False, "center")
            else:
                c = t.cell(ri + 1, 0)
                if v not in (None, ""):
                    _put(c, v, size, ri in bold_rows, "center")
            ri = rj + 1
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(text), 12.5, True)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "8")
    bt.set(qn("w:color"), "000000")
    bdr.append(bt)
    pPr.append(bdr)


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(text), 10.5, True)


def para(doc, text, size=9.5, indent=0, color=None, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.35
    if indent:
        p.paragraph_format.left_indent = Mm(indent)
    _font(p.add_run(text), size, False, color)
    return p


def note(doc, text):
    return para(doc, "※ " + text, 8, color="404040", after=4)


def _txt(v):
    return v.strip() if isinstance(v, str) and v.strip() else "—"


def _page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        _font(r, 8)
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)


# ───────────────────────── 본문 ─────────────────────────
def cover(doc, data):
    pr = data["profile"]
    a, b = core.periods(data)["month"]
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("개 인 과 제   보 고 서"), 11, True, "404040")
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t, outer=18, inner=0)
    _widths(t, [150])
    c = t.rows[0].cells[0]
    _put(c, f"{a.year}년 {a.month}월 개인 현금흐름 관리 보고서", 18, True, "center")
    q = c.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(8)
    _font(q.add_run(pr.get("title") or ""), 11)
    for _ in range(8):
        doc.add_paragraph()
    info = [("교 과 목", pr.get("course")), ("소    속", pr.get("dept")), ("학    번", pr.get("student_id")),
            ("성    명", pr.get("name")), ("기록 기간", f"{a:%Y.%m.%d} ~ {b:%Y.%m.%d}"),
            ("제 출 일", pr.get("submit_date") or date.today().strftime("%Y.%m.%d"))]
    t2 = doc.add_table(rows=0, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t2)
    for k, v in info:
        cells = t2.add_row().cells
        _put(cells[0], k, 10, True, "center")
        _shade(cells[0], "F2F2F2")
        _put(cells[1], _txt(v), 10, False, "left")
    _widths(t2, [35, 75])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc(doc):
    h1(doc, "목  차")
    items = ["Ⅰ. 개요 및 관리 범위", "Ⅱ. 개인 재무상태표", "Ⅲ. 한 달 예산 수립", "Ⅳ. 2주 기간 예산",
             "Ⅴ. 현금흐름 기록", "Ⅵ. 결산 및 현금흐름 분석", "Ⅶ. 분석의 한계", "Ⅷ. 절약 목표와 실행 방안",
             "Ⅸ. 느낀 점과 현금흐름 관리 원칙", "부록. 지출 증빙"]
    for i in items:
        para(doc, i, 10, indent=5, after=4)
    note(doc, "별첨: 부속명세서(예산표, 예산 대비 실적, 일별 잔액, 거래 원장, 카테고리 분석)는 "
              "본 보고서 뒤에 Excel 원본에서 출력하여 첨부함.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def sec_overview(doc, data):
    pr = data["profile"]
    ps = core.periods(data)
    h1(doc, "Ⅰ. 개요 및 관리 범위")
    hol = ", ".join(f"{core.d(x).month}/{core.d(x).day}" for x in data["holidays"]) or "—"
    rows = [["생활 형태", _txt(pr.get("life_type"))],
            ["주요 소득", _txt(pr.get("income_sources"))],
            ["관리 대상", _txt(pr.get("managed"))],
            ["비용 구분", _txt(pr.get("cost_split"))],
            ["예산 대상 월", f"{ps['month'][0].year}년 {ps['month'][0].month}월"],
            ["기록 기간", f"{core.period_title(data, 'p1')}, {core.period_title(data, 'p2')}"],
            ["휴일 반영", f"주말 및 공휴일({hol})"]]
    table(doc, ["구분", "내용"], rows, [35, 135], center=(0,))
    h2(doc, "가족 등 타인 부담 비용 (본인 지출에서 제외, 별도 기록)")
    fp = data["family_paid"]
    if fp:
        rows = [[x.get("name"), core.parse_amount(x.get("amount")), x.get("payer"), x.get("note")] for x in fp]
        rows.append(["합계", sum(core.parse_amount(x.get("amount")) for x in fp), "", ""])
        table(doc, ["항목", "월 금액(원)", "부담자", "비고"], rows, [45, 30, 30, 65], num=(1,), center=(2,),
              bold_rows=(len(rows) - 1,), shade_rows=(len(rows) - 1,))
    else:
        para(doc, "해당 없음")


def sec_balance(doc, data):
    h1(doc, "Ⅱ. 개인 재무상태표")
    bs = core.balance_sheet(data)
    a, b = core.periods(data)["month"]
    rows, bold = [], []
    tot = {}
    for typ in core.ASSET_TYPES:
        items = [x for x in bs["items"] if x["type"] == typ]
        if not items:
            continue
        for i, x in enumerate(items):
            e = x["end"]
            rows.append([typ, x["name"], x["start"], e if e is not None else "미입력",
                         (e - x["start"]) if e is not None else "—"])
        s = sum(x["start"] for x in items)
        ok = all(x["end"] is not None for x in items)
        e = sum(x["end"] for x in items) if ok else None
        tot[typ] = (s, e)
        rows.append(["", f"{typ} 합계", s, e if e is not None else "미입력", (e - s) if e is not None else "—"])
        bold.append(len(rows) - 1)
    s_net = sum(v[0] for k, v in tot.items() if k != "부채") - tot.get("부채", (0, 0))[0]
    ends = [v[1] for v in tot.values()]
    e_net = (sum(v[1] for k, v in tot.items() if k != "부채") - (tot.get("부채", (0, 0))[1] or 0)) \
        if all(x is not None for x in ends) and ends else None
    rows.append(["순자산", "자산 − 부채", s_net, e_net if e_net is not None else "미입력",
                 (e_net - s_net) if e_net is not None else "—"])
    bold.append(len(rows) - 1)
    table(doc, ["구분", "계정", f"기초({a:%m/%d})", f"기말({b:%m/%d})", "증감"], rows, [26, 52, 30, 30, 32],
          num=(2, 3, 4), bold_rows=bold, shade_rows=bold, merge_first=True)
    note(doc, f"기간 중 저축·투자 납입액 {core.won(bs['savings_added'])}원, "
              f"기록 기준 기말 가용 현금(장부) {core.won(core.analyze(data, 'month')['book_end'])}원. "
              "기말 금액은 실제 계좌·현금 잔액을 확인하여 입력함.")


def sec_budget(doc, data):
    h1(doc, "Ⅲ. 한 달 예산 수립")
    h2(doc, "1. 가용 예산 산출 (3개 안 비교)")
    bo = core.budget_options(data)
    o = bo["options"]
    ch = data["budget"].get("choice")
    heads = ["구분"] + [f"{'■ ' if k == ch else ''}{v['name']}({k})" for k, v in o.items()]
    R = lambda label, f: [label] + [f(v) for v in o.values()]
    rows = [
        R("기초 가용 잔액", lambda v: bo["start"]),
        R("(+) 예상 수입", lambda v: bo["income"]),
        R("(−) 저축 계획", lambda v: bo["savings"]),
        R("(−) 고정지출", lambda v: bo["fixed"]),
        R("(+) 이월 잔액 활용", lambda v: v["carry_used"]),
        R("가용 예산", lambda v: v["available"]),
        R("  예비비", lambda v: v["reserve"]),
        R("  변동지출 배정", lambda v: v["variable"]),
        R("월말 예상 이월액", lambda v: v["carry_end"]),
        R("일 평균 변동지출", lambda v: int(v["daily"])),
        R("저축률(저축/수입)", lambda v: f"{v['saving_rate']:.1f}%"),
        R("산출 기준", lambda v: v["desc"]),
    ]
    table(doc, heads, rows, [38, 44, 44, 44], num=(1, 2, 3), bold_rows=(5,), shade_rows=(5,), size=8)
    note(doc, f"선택안: {core.TIERS[ch]['name'] if ch else '미선택'}. "
              "가용 예산 = 예상 수입 − 저축 − 고정지출 + 이월 잔액 활용액, 예비비는 가용 예산의 10%(천원 단위 절사).")
    h2(doc, "2. 한 달 예산표")
    rows = [[r[0], r[1], r[2], r[5]] for r in core.budget_rows(data)]
    table(doc, ["구분", "항목", "한 달 계획 금액(원)", "설정 근거"], rows, [24, 40, 32, 74], num=(2,),
          merge_first=True)
    s = core.budget_summary(data)
    table(doc, ["시작 잔액", "기대 소득", "저축 계획", "지출 배정(고정+변동)", "예비비"],
          [[s["start"], s["income"], s["savings"], s["fixed"] + s["variable"], s["reserve"]]],
          [34, 34, 34, 34, 34], num=(0, 1, 2, 3, 4))


def sec_split(doc, data):
    h1(doc, "Ⅳ. 2주 기간 예산")
    _, wd1, h1_ = core.day_profile(data, "p1")
    n2, wd2, h2_ = core.day_profile(data, "p2")
    para(doc, f"{core.period_title(data, 'p1')}: 평일 {wd1}일·휴일 {h1_}일 / "
              f"{core.period_title(data, 'p2')}: 평일 {wd2}일·휴일 {h2_}일")
    rows = [[r[0], r[1], r[2], r[3], r[4], r[6]] for r in core.budget_rows(data)]
    rows.append(["", "합계(소득 제외)", *[sum(r[i] for r in core.budget_rows(data) if r[0] != "소득")
                                       for i in (2, 3, 4)], ""])
    table(doc, ["구분", "항목", "한 달", "상반기", "후반기", "2주 설정 근거"], rows, [20, 30, 22, 22, 22, 54],
          num=(2, 3, 4), merge_first=True, bold_rows=(len(rows) - 1,), shade_rows=(len(rows) - 1,), size=8)
    note(doc, "월 예산을 절반으로 나누지 않음. 수입·저축·고정지출은 예정일이 속한 기간에 전액 배정하고, "
              "변동지출은 예정 일정 금액을 먼저 해당 기간에 배정한 뒤 나머지를 카테고리별 평일·휴일 가중 일수로 배분함.")


def sec_records(doc, data):
    h1(doc, "Ⅴ. 현금흐름 기록")
    for key in ("p1", "p2"):
        an = core.analyze(data, key)
        h2(doc, f"{an['title']} 기록 — {len(an['txs'])}건")
        rows = []
        for t in an["txs"]:
            dt = core.d(t["date"])
            cat = t["category"] if t["kind"] == "지출" else t["kind"]
            rows.append([f"{dt.month}/{dt.day}", t["item"], t["amount"], cat, t["method"],
                         "계획" if t["planned"] else "비계획", _txt(t["reason"]), _txt(t["emotion"])])
        if rows:
            table(doc, ["날짜", "지출 항목", "금액(원)", "카테고리", "결제수단", "계획 여부", "이유·상황", "감정"],
                  rows, [11, 25, 17, 18, 16, 13, 33, 37], num=(2,), center=(0, 3, 4, 5), size=7.5)
        ns = ", ".join(f"{x.month}/{x.day}" for x in an["no_spend"])
        para(doc, f"지출 없음: {ns if ns else '해당 없음'}", 8.5)


def _variance(doc, data, an):
    rows, bold = [], []
    for g, n, b, a in an["rows"]:
        rows.append([g, n, b, a, a - b, core.remark(data, an, g, n)])
    for label, b, a in (("유입 합계", an["inflow_b"], an["inflow_a"]), ("유출 합계", an["out_b"], an["out_a"]),
                        ("순현금흐름(유입−유출)", an["net_b"], an["net_a"])):
        rows.append(["", label, b, a, a - b, ""])
        bold.append(len(rows) - 1)
    rows.append(["", "예비비", an["reserve_b"], an["reserve_used"], an["reserve_b"] - an["reserve_used"],
                 "실적 = 예산 초과 소비분, 차이 = 잔여 예비비"])
    bold.append(len(rows) - 1)
    for r in rows:
        r[4] = f"{r[4]:+,}" if r[4] else "0"
    table(doc, ["구분", "항목", "기간 예산", "실제 흐름", "차이", "비고"], rows, [18, 30, 21, 21, 20, 60],
          num=(2, 3, 4), bold_rows=bold, shade_rows=bold, merge_first=True, size=8)


def _facts(an):
    out = []
    for g, n, b, a in an["rows"]:
        if g != "변동지출" or b == a:
            continue
        ts = [t for t in an["txs"] if t["kind"] == "지출" and t["category"] == n]
        s = f"{n}: 기간 예산 {core.won(b)}원 대비 실제 {core.won(a)}원으로 {core.won(abs(a - b))}원 {'초과' if a > b else '미달'}"
        if ts:
            subs = {}
            for t in ts:
                subs[t["sub"]] = subs.get(t["sub"], 0) + t["amount"]
            top = max(subs, key=subs.get)
            up = [t for t in ts if not t["planned"]]
            s += f"(기록 {len(ts)}건, 세부 항목 중 {top} {core.won(subs[top])}원"
            s += f", 비계획 {len(up)}건 {core.won(sum(t['amount'] for t in up))}원)" if up else ")"
        out.append(s)
    return out


def sec_settle(doc, data):
    h1(doc, "Ⅵ. 결산 및 현금흐름 분석")
    ans = {k: core.analyze(data, k) for k in ("p1", "p2", "month")}
    n = 1
    for key in ("p1", "p2"):
        an = ans[key]
        h2(doc, f"{n}. {an['title']} 예산 대비 실적")
        _variance(doc, data, an)
        doc.add_picture(charts.budget_vs_actual(an), width=Mm(150))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        h2(doc, "차이 해석")
        for f in _facts(an):
            para(doc, "· " + f, 9, indent=2, after=1)
        it = data["settle"].get(key, {}).get("interpret", "")
        if it.strip():
            para(doc, it, 9.5, after=4)
        n += 1
    m = ans["month"]
    h2(doc, f"{n}. 월간 현금흐름표")
    n += 1
    rows, bold = [["Ⅰ. 기초 현금", "", m["start"]]], [0]
    rows.append(["Ⅱ. 현금 유입", "", m["income"]])
    bold.append(len(rows) - 1)
    for g, nm, b, a in m["rows"]:
        if g == "소득" and nm != "기초 잔액(이월)" and a:
            rows.append(["", nm, a])
    rows.append(["Ⅲ. 현금 유출", "", m["expense"] + m["saving"]])
    bold.append(len(rows) - 1)
    fixed = sum(a for g, _, _, a in m["rows"] if g == "고정지출")
    rows.append(["  1. 소비지출", "", m["expense"]])
    rows.append(["", "고정지출", fixed])
    for g, nm, b, a in m["rows"]:
        if g == "변동지출" and a:
            rows.append(["", nm, a])
    rows.append(["  2. 저축·투자", "", m["saving"]])
    for g, nm, b, a in m["rows"]:
        if g == "저축" and a:
            rows.append(["", nm, a])
    rows.append(["Ⅳ. 순현금흐름(Ⅱ−Ⅲ)", "", m["income"] - m["expense"] - m["saving"]])
    bold.append(len(rows) - 1)
    rows.append(["Ⅴ. 기말 현금(장부, Ⅰ+Ⅳ)", "", m["book_end"]])
    bold.append(len(rows) - 1)
    ae = m["actual_end"]
    rows.append(["Ⅵ. 기말 현금(실제 확인)", "", ae if ae is not None else "미입력"])
    rows.append(["   장부 − 실제 차이", "", (m["book_end"] - ae) if ae is not None else "—"])
    table(doc, ["과목", "세부", "금액(원)"], rows, [70, 55, 45], num=(2,), bold_rows=bold, shade_rows=bold)

    h2(doc, f"{n}. 카테고리별 지출 분석")
    n += 1
    cats = m["cats"]
    order = ["고정지출"] + core.VAR_CATS
    bud = {r[1]: r[2] for r in core.budget_rows(data) if r[0] == "변동지출"}
    bud["고정지출"] = sum(r[2] for r in core.budget_rows(data) if r[0] == "고정지출")
    rows = []
    for c in order:
        if c not in cats:
            continue
        x = cats[c]
        p1 = ans["p1"]["cats"].get(c, {}).get("amount", 0)
        p2 = ans["p2"]["cats"].get(c, {}).get("amount", 0)
        rows.append([c, p1, p2, x["amount"], x["count"], f"{x['pct_exp']:.1f}%", f"{x['pct_inc']:.1f}%",
                     f"{x['amount'] / bud[c] * 100:.0f}%" if bud.get(c) else "—"])
    rows.append(["합계", ans["p1"]["expense"], ans["p2"]["expense"], m["expense"], m["count"], "100.0%",
                 f"{m['expense'] / m['income'] * 100:.1f}%" if m["income"] else "—", ""])
    table(doc, ["카테고리", "상반기", "후반기", "월 합계", "횟수", "지출 대비", "소득 대비", "예산 대비"], rows,
          [28, 22, 22, 24, 14, 20, 20, 20], num=(1, 2, 3, 4), center=(5, 6, 7),
          bold_rows=(len(rows) - 1,), shade_rows=(len(rows) - 1,), size=8)
    note(doc, "소득 대비 = 기간 중 실제 유입 수입(기초 잔액 제외) 대비 비중.")

    h2(doc, f"{n}. 소비 패턴")
    n += 1
    wdd, hdd = m["weekday"], m["holiday"]
    tot = m["expense"] or 1
    rows = [["평일", wdd["days"], wdd["amount"], int(wdd["amount"] / max(wdd["days"], 1)),
             f"{wdd['amount'] / tot * 100:.1f}%"],
            ["휴일(주말·공휴일)", hdd["days"], hdd["amount"], int(hdd["amount"] / max(hdd["days"], 1)),
             f"{hdd['amount'] / tot * 100:.1f}%"]]
    table(doc, ["구분", "일수", "지출액", "일 평균", "비중"], rows, [50, 20, 35, 35, 30], num=(1, 2, 3), center=(4,))
    rows = [["계획 지출", m["planned"], f"{m['planned'] / tot * 100:.1f}%"],
            ["비계획 지출", m["unplanned"], f"{m['unplanned'] / tot * 100:.1f}%"],
            ["지출 없는 날", f"{len(m['no_spend'])}일", ""]]
    table(doc, ["구분", "금액·일수", "비중"], rows, [60, 55, 55], num=(1,), center=(2,))
    rows = [[f"반복 소액: {c} / {s}", f"{k}회", amt] for c, s, k, amt in m["small"]]
    rows += [[f"일회성 대형: {core.d(t['date']).month}/{core.d(t['date']).day} {t['item']}", "1회", t["amount"]]
             for t in m["big"]]
    if rows:
        table(doc, ["구분 / 항목", "횟수", "금액(원)"], rows, [100, 25, 45], num=(2,), center=(1,))
    note(doc, f"반복 소액 = 1만원 미만이며 같은 세부 항목 3회 이상, 일회성 대형 = 단건 {core.won(m['big_th'])}원 이상"
              "(3만원과 월 변동지출 예산 10% 중 큰 값, 고정지출 제외).")

    h2(doc, f"{n}. 유지·조정·중단 분류")
    n += 1
    rows = []
    for j, desc in (("유지", "필요하거나 만족도가 높은 지출"), ("조정", "목적은 필요하나 가격·빈도 조정 가능"),
                    ("중단", "만족도가 낮은 비계획 지출")):
        ts = m["judged"][j]
        items = {}
        for t in ts:
            items[t["item"]] = items.get(t["item"], 0) + t["amount"]
        top = ", ".join(k for k, _ in sorted(items.items(), key=lambda x: -x[1])[:4])
        rows.append([j, desc, len(ts), sum(t["amount"] for t in ts), top or "—"])
    table(doc, ["분류", "기준", "건수", "금액(원)", "주요 항목"], rows, [14, 50, 14, 24, 68], num=(2, 3), center=(0,))
    note(doc, "분류 기준: 만족도 4 이상 → 유지, 만족도 2 이하이면서 비계획 → 중단, 그 외 비계획·대체 가능 항목 → 조정. "
              "개별 거래는 직접 판단으로 변경 가능.")

    h2(doc, f"{n}. 잔액 점검")
    inc_days = [core.d(i["date"]) for i in data["incomes"] if i.get("date")]
    pre = [x for x in m["daily"] if inc_days and x["date"] < max(inc_days)]
    pre_min = min(pre, key=lambda x: x["bal"]) if pre else None
    fixed_next = sum(core.parse_amount(x.get("amount")) for x in data["fixed"])
    rows = [["기록 시작 잔액", core.won(m["start"]) + "원"],
            ["기간 중 유입 / 유출", f"{core.won(m['income'])}원 / {core.won(m['expense'] + m['saving'])}원"],
            ["기록으로 계산한 종료 잔액", core.won(m["book_end"]) + "원"],
            ["실제 종료 잔액", (core.won(m["actual_end"]) + "원") if m["actual_end"] is not None else "미입력"],
            ["일치 여부", "일치" if m["actual_end"] == m["book_end"] else
             (f"불일치({m['book_end'] - m['actual_end']:+,}원)" if m["actual_end"] is not None else "—")],
            ["기간 중 최저 잔액", f"{core.won(m['min_bal'])}원 ({m['min_day'].month}/{m['min_day'].day})"],
            ["수입 입금 전 최저 잔액", f"{core.won(pre_min['bal'])}원 ({pre_min['date'].month}/{pre_min['date'].day})"
             if pre_min else "—"],
            ["다음 달 고정지출 감당", f"고정지출 {core.won(fixed_next)}원 대비 종료 잔액 "
                                f"{(m['book_end'] / fixed_next):.1f}배" if fixed_next else "—"],
            ["잔여 예비비", f"{core.won(m['reserve_b'] - m['reserve_used'])}원"]]
    table(doc, ["점검 항목", "결과"], rows, [60, 110], center=())
    doc.add_picture(charts.daily_balance(m), width=Mm(150))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def sec_limits(doc, data):
    h1(doc, "Ⅶ. 분석의 한계")
    ev = data["evaluation"]
    rows = [["■" if x in ev["limits"] else "□", x] for x in core.LIMIT_TYPES]
    table(doc, ["해당", "한계 유형"], rows, [15, 155], center=(0,))
    if ev.get("limit_text", "").strip():
        para(doc, ev["limit_text"])


def sec_goals(doc, data):
    h1(doc, "Ⅷ. 절약 목표와 실행 방안")
    goals = data["evaluation"]["goals"]
    if not goals:
        para(doc, "—")
    labels = [("problem", "발견한 문제"), ("target", "목표"), ("action", "행동"), ("saving", "예상 조정액"),
              ("check", "점검 방법"), ("use", "절약액의 용도")]
    for i, g in enumerate(goals, 1):
        h2(doc, f"목표 {i}")
        table(doc, ["구성 요소", "내용"], [[l, _txt(g.get(k))] for k, l in labels], [35, 135], center=(0,))


def sec_principles(doc, data):
    h1(doc, "Ⅸ. 느낀 점과 현금흐름 관리 원칙")
    ev = data["evaluation"]
    h2(doc, "1. 느낀 점")
    para(doc, _txt(ev.get("feelings")))
    h2(doc, "2. 나의 현금흐름 관리 원칙")
    ps = [p for p in ev.get("principles", []) if (p.get("principle") or "").strip()]
    if ps:
        table(doc, ["No.", "원칙", "실천 확인 방법"],
              [[i, p.get("principle"), _txt(p.get("check"))] for i, p in enumerate(ps, 1)],
              [12, 90, 68], center=(0,))
    else:
        para(doc, "—")


def sec_receipts(doc, data):
    rs = [t for t in sorted(data["transactions"], key=lambda t: t["date"])
          if t.get("receipt") and Path(t["receipt"]).suffix.lower() in (".png", ".jpg", ".jpeg", ".bmp", ".gif")
          and Path(t["receipt"]).exists()]
    if not rs:
        return
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h1(doc, "부록. 지출 증빙")
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)
    for i in range(0, len(rs), 2):
        cells = t.add_row().cells
        _no_split(t.rows[-1])
        for j, tx in enumerate(rs[i:i + 2]):
            c = cells[j]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p.add_run().add_picture(tx["receipt"], height=Mm(80))
            except Exception:
                _font(p.add_run("(이미지를 열 수 없음)"), 8)
            q = c.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            dt = core.d(tx["date"])
            _font(q.add_run(f"{dt.month}/{dt.day} {tx['item']} {core.won(tx['amount'])}원"), 8)
    _widths(t, [85, 85])


def build(data, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(9.5)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
    _page_number(sec)
    sec.different_first_page_header_footer = True
    cover(doc, data)
    toc(doc)
    sec_overview(doc, data)
    sec_balance(doc, data)
    sec_budget(doc, data)
    sec_split(doc, data)
    sec_records(doc, data)
    sec_settle(doc, data)
    sec_limits(doc, data)
    sec_goals(doc, data)
    sec_principles(doc, data)
    sec_receipts(doc, data)
    doc.save(path)
    return path
